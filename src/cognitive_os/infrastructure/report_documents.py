"""Deterministic, source-linked report exports. No additional model request."""

from datetime import UTC, datetime
from pathlib import Path
import re
from xml.sax.saxutils import escape

from cognitive_os.schemas.recordings import VisualReportContent

MEDIA_TYPES = {"pdf": "application/pdf",
               "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
SOURCES = {"notes": "Notas", "transcript": "Transcripcion", "clarifications": "Aclaraciones"}


def clean(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(value))


def timestamp(milliseconds):
    seconds = milliseconds // 1000
    return f"{seconds // 60:02d}:{seconds % 60:02d}.{milliseconds % 1000:03d}"


def document_blocks(snapshot, frames, style):
    report = VisualReportContent.model_validate(snapshot["content"])
    blocks = []

    def add(kind, text):
        blocks.append((kind, clean(text)))

    def sources(statement):
        labels = [SOURCES.get(source, source) for source in statement.text_sources]
        labels += [f"Video {timestamp(snapshot['sampling']['frames'][index]['timestamp_ms'])}"
                   for index in statement.frame_indices]
        if labels:
            add("caption", "Fuentes: " + "; ".join(labels))

    add("title", report.title)
    add("subtitle", "Tutorial paso a paso" if style == "tutorial" else "Informe de aprendizaje")
    status = "APROBADO" if snapshot["review_status"] == "approved" else "BORRADOR - pendiente de validacion"
    add("status", f"{status} | Revision {snapshot['revision']}")
    add("caption", "Exportado: " + datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC"))
    if snapshot.get("reviewed_at"):
        add("caption", "Revision humana: " + snapshot["reviewed_at"])
    add("heading", "Resumen")
    add("body", report.summary)
    if style == "report":
        add("heading", "Informe")
        add("body", report.report)
    for key, title in (("prerequisites", "Antes de comenzar"), ("business_rules", "Reglas del proceso"),
                       ("exceptions", "Excepciones")):
        items = getattr(report, key)
        if items:
            add("heading", title)
            for item in items:
                add("body", item.text)
                sources(item)
    add("heading", "Actividades del usuario")
    for number, step in enumerate(report.instructions, 1):
        add("step", f"Paso {number}")
        add("body", step.instruction)
        add("body", "Resultado esperado: " + step.expected_result)
        sources(step)
        for alternative in step.alternatives:
            destination = f"paso {alternative.target_step}" if alternative.target_step else "fin"
            add("body", f"Si {alternative.condition}: continuar en {destination}.")
        # First and last cited captures bound document size without inventing illustrations.
        indices = list(dict.fromkeys(step.frame_indices))
        selected = indices[:1] + indices[-1:] if len(indices) > 1 else indices
        for index in selected:
            image, time_ms = frames[index]
            blocks.append(("image", image))
            add("caption", f"Paso {number} - captura {index + 1} - video {timestamp(time_ms)}")
        if not indices:
            add("caption", "Actividad sustentada en fuentes textuales; sin captura visual citada.")
    if snapshot["clarifications"]:
        add("heading", "Aclaraciones resueltas")
        for item in snapshot["clarifications"]:
            add("step", item["question"])
            add("body", item["answer"])
    if report.uncertainties:
        add("heading", "Limitaciones y puntos por verificar")
        for uncertainty in report.uncertainties:
            add("body", uncertainty)
    add("heading", "Trazabilidad")
    add("caption", f"Video: {snapshot['recording_id']} | Modelo: {snapshot['model_name']}")
    add("body", "Las capturas corresponden a fotogramas muestreados del video original. "
        "No representan una observacion continua ni garantizan que se hayan identificado todas las acciones. "
        "Las aclaraciones se adjuntan tal como fueron respondidas; para incorporarlas en los pasos, "
        "regenera y revisa el informe antes de aprobarlo.")
    return blocks


def build_document(snapshot, frames, target: Path, *, style="tutorial", format="pdf"):
    blocks = document_blocks(snapshot, frames, style)
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from PIL import Image

        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.8)
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
            font = document.styles[name].font
            font.name, font.color.rgb = "Arial", RGBColor.from_string("20262C")
        document.styles["Normal"].font.size = Pt(10)
        document.styles["Normal"].paragraph_format.space_after = Pt(7)
        document.styles["Title"].font.size = Pt(24)
        document.styles["Heading 1"].font.size = Pt(16)
        document.styles["Heading 2"].font.size = Pt(12)
        styles = {"title": "Title", "subtitle": "Subtitle", "heading": "Heading 1",
                  "step": "Heading 2", "caption": "Caption"}
        for kind, value in blocks:
            if kind == "image":
                with Image.open(value) as image:
                    width = min(6.5, 4.5 * image.width / image.height)
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(str(value), width=Inches(width))
            else:
                paragraph = document.add_paragraph(value, style=styles.get(kind, "Normal"))
                if kind == "status":
                    paragraph.runs[0].bold = True
        document.core_properties.author = "Cognitive OS"
        document.core_properties.title = clean(snapshot["content"]["title"])
        document.save(target)
    elif format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

        styles = {
            "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=15, spaceAfter=8),
            "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=24, leading=29, spaceAfter=12),
            "subtitle": ParagraphStyle("subtitle", fontSize=13, leading=18, spaceAfter=10),
            "heading": ParagraphStyle("heading", fontName="Helvetica-Bold", fontSize=16, leading=21,
                                      spaceBefore=16, spaceAfter=8, keepWithNext=True),
            "step": ParagraphStyle("step", fontName="Helvetica-Bold", fontSize=12, leading=17,
                                   spaceBefore=12, spaceAfter=6, keepWithNext=True),
            "caption": ParagraphStyle("caption", fontSize=8, leading=12, textColor=colors.HexColor("#4b5563"),
                                      spaceAfter=10),
            "status": ParagraphStyle("status", fontName="Helvetica-Bold", fontSize=10, leading=15, spaceAfter=10),
        }
        story = []
        pending_image = None
        for kind, value in blocks:
            if kind == "image":
                image = Image(str(value))
                scale = min(470 / image.imageWidth, 320 / image.imageHeight)
                image.drawWidth, image.drawHeight = image.imageWidth * scale, image.imageHeight * scale
                pending_image = image
            else:
                paragraph = Paragraph(escape(value).replace("\n", "<br/>"), styles[kind])
                if pending_image is not None:
                    story.append(KeepTogether([pending_image, Spacer(1, 4), paragraph]))
                    pending_image = None
                else:
                    story.append(paragraph)
        SimpleDocTemplate(str(target), pagesize=A4, rightMargin=57, leftMargin=57,
                          topMargin=54, bottomMargin=54, title=clean(snapshot["content"]["title"]),
                          author="Cognitive OS").build(story)
    else:
        raise ValueError("Unsupported document format")
